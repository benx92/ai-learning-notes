"""将本地文件解析为带原始位置的文本单元，不负责切分或向量化。"""

import csv
import json
from dataclasses import dataclass
from pathlib import Path


@dataclass
class DocumentUnit:
    text: str
    metadata: dict


class DocumentLoader:
    def load(self, file_path):
        path = Path(file_path).resolve()
        if not path.is_file():
            raise ValueError("请指定一个存在的文件，不支持目录入库")
        readers = {".md": self._plain, ".txt": self._plain,
                   ".pdf": self._pdf, ".docx": self._docx,
                   ".pptx": self._pptx, ".xlsx": self._xlsx,
                   ".html": self._html, ".htm": self._html,
                   ".csv": self._csv, ".json": self._json,
                   ".jsonl": self._jsonl}
        reader = readers.get(path.suffix.lower())
        if reader is None:
            raise ValueError("不支持此格式；旧版 .doc/.ppt/.xls 请先转换为新版格式")
        units = []
        for text, metadata in reader(path):
            text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
            if text:
                units.append(DocumentUnit(text, {"source": str(path),
                    "file_type": path.suffix.lower(), **metadata}))
        if not units:
            raise ValueError("文件没有可用正文，保留原有索引")
        return units

    @staticmethod
    def _plain(path):
        yield path.read_text(encoding="utf-8-sig"), {"location": "全文"}

    @staticmethod
    def _pdf(path):
        import pymupdf

        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ValueError("PDF 已加密，请先解密后再解析")
            for number, page in enumerate(document, 1):
                text = page.get_text("text", sort=True).strip()
                if not text:
                    raise ValueError(f"PDF 第 {number} 页未提取到文字；请检查空白页或先做 OCR")
                yield text, {"page": number, "location": f"PDF 第 {number} 页"}

    @staticmethod
    def _docx(path):
        from docx import Document
        from docx.table import Table

        for block_index, block in enumerate(Document(path).iter_inner_content(), 1):
            if isinstance(block, Table):
                for row_number, row in enumerate(block.rows, 1):
                    text = " | ".join(cell.text.strip() for cell in row.cells)
                    if any(cell.text.strip() for cell in row.cells):
                        yield text, {"block": block_index, "row": row_number,
                            "location": f"Word 块 {block_index}，表格第 {row_number} 行"}
            else:
                yield block.text, {"block": block_index,
                                   "location": f"Word 块 {block_index}"}

    @staticmethod
    def _pptx(path):
        from pptx import Presentation

        def shape_texts(shapes):
            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)
                if hasattr(shape, "shapes"):
                    yield from shape_texts(shape.shapes)

        for number, slide in enumerate(Presentation(path).slides, 1):
            yield "\n".join(shape_texts(slide.shapes)), {
                "slide": number, "location": f"PPT 第 {number} 张幻灯片"}

    @staticmethod
    def _xlsx(path):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                headers = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ["" if value is None else str(value) for value in row]
                    if not any(values):
                        continue
                    if headers is None:
                        headers = [value or f"列{i}" for i, value in enumerate(values, 1)]
                        text = "表头：" + " | ".join(headers)
                    else:
                        text = "；".join(f"{headers[i] if i < len(headers) else f'列{i+1}'}={value}"
                                        for i, value in enumerate(values) if value)
                    yield f"工作表：{sheet.title}\n{text}", {
                        "sheet": sheet.title, "row": number,
                        "location": f"Excel 工作表 {sheet.title}，第 {number} 行"}
        finally:
            workbook.close()

    @staticmethod
    def _html(path):
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8-sig"), "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        yield soup.get_text(separator="\n", strip=True), {"location": "HTML 正文"}

    @staticmethod
    def _csv(path):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            for number, row in enumerate(csv.reader(stream), 1):
                if any(row):
                    yield json.dumps(row, ensure_ascii=False), {
                        "row": number, "location": f"CSV 第 {number} 条记录"}

    @staticmethod
    def _json(path):
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        yield json.dumps(value, ensure_ascii=False, indent=2), {"location": "JSON 全文"}

    @staticmethod
    def _jsonl(path):
        for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            if line.strip():
                yield json.dumps(json.loads(line), ensure_ascii=False), {
                    "row": number, "location": f"JSONL 第 {number} 行"}
