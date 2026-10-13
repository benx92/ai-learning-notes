"""真实文件解析与来源传递检查；向量和生成服务仍使用测试替身。"""

import tempfile
import unittest
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

import rag
from loader import DocumentLoader
from test_rag import TestEmbedder


class LoaderChecks(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.loader = DocumentLoader()

    def test_pdf_pages_and_empty_page(self):
        units = self.loader.load(Path(__file__).with_name("refund-policy.pdf"))
        self.assertEqual([unit.metadata["page"] for unit in units], [1, 2])
        self.assertIn("15 日", units[0].text)
        self.assertIn("定制商品", units[1].text)
        empty = self.root / "empty.pdf"
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(empty)
        with self.assertRaisesRegex(ValueError, "OCR"):
            self.loader.load(empty)

    def test_word_block_order_and_table(self):
        path = self.root / "policy.docx"
        doc = Document()
        doc.add_paragraph("表格前的规则")
        row = doc.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = "会员", "15日"
        doc.add_paragraph("表格后的例外")
        doc.save(path)
        units = self.loader.load(path)
        self.assertEqual([unit.metadata["block"] for unit in units], [1, 2, 3])
        self.assertEqual(units[1].text, "会员 | 15日")
        self.assertEqual(units[1].metadata["row"], 1)

    def test_ppt_slide_and_table(self):
        path = self.root / "policy.pptx"
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "会员退货"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "会员", "15日"
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "定制商品例外"
        ppt.save(path)
        units = self.loader.load(path)
        self.assertEqual([unit.metadata["slide"] for unit in units], [1, 2])
        self.assertIn("会员 | 15日", units[0].text)

    def test_excel_header_sheet_row_and_zero(self):
        path = self.root / "policy.xlsx"
        wb = Workbook()
        sheet = wb.active
        sheet.title = "售后规则"
        sheet.append(["类别", "期限"])
        sheet.append(["会员", 15])
        sheet.append([None, None])
        sheet.append(["手续费", 0])
        wb.save(path)
        wb.close()
        units = self.loader.load(path)
        self.assertEqual([unit.metadata["row"] for unit in units], [1, 2, 4])
        self.assertEqual(units[1].metadata["sheet"], "售后规则")
        self.assertIn("类别=会员；期限=15", units[1].text)
        self.assertIn("期限=0", units[2].text)

    def test_html_csv_json_and_jsonl(self):
        html = self.root / "policy.html"
        html.write_text('<h1>会员规则</h1><script>噪声</script><p>15日</p>', encoding="utf-8")
        self.assertEqual(self.loader.load(html)[0].text, "会员规则\n15日")
        csv = self.root / "policy.csv"
        csv.write_text('类别,说明\n会员,"15日,可申请"\n', encoding="utf-8")
        units = self.loader.load(csv)
        self.assertIn('15日,可申请', units[1].text)
        self.assertEqual(units[1].metadata["row"], 2)
        json = self.root / "policy.json"
        json.write_text('{"期限":15}', encoding="utf-8")
        self.assertIn('"期限": 15', self.loader.load(json)[0].text)
        jsonl = self.root / "policy.jsonl"
        jsonl.write_text('{"期限":15}\n\n{"期限":7}', encoding="utf-8")
        self.assertEqual([u.metadata["row"] for u in self.loader.load(jsonl)], [1, 3])
        jsonl.write_text('not json', encoding="utf-8")
        with self.assertRaises(ValueError):
            self.loader.load(jsonl)

    def test_reject_legacy_format_and_invalid_utf8(self):
        old = self.root / "old.doc"
        old.write_bytes(b"legacy")
        with self.assertRaisesRegex(ValueError, "转换"):
            self.loader.load(old)
        bad = self.root / "bad.txt"
        bad.write_bytes(b"\xff")
        with self.assertRaises(UnicodeError):
            self.loader.load(bad)

    def test_pdf_location_survives_index_context_and_update_failure(self):
        path = self.root / "policy.pdf"
        path.write_bytes(Path(__file__).with_name("refund-policy.pdf").read_bytes())
        pipeline = rag.Pipeline(self.root / "db", TestEmbedder())
        result = pipeline.ingest(path)
        self.assertEqual(result["units"], 2)
        initial = set(pipeline.collection.get(include=[])["ids"])
        pipeline.ingest(path)
        self.assertEqual(initial, set(pipeline.collection.get(include=[])["ids"]))
        hits = pipeline.retrieve("会员", top_k=100)
        self.assertEqual({hit["metadata"]["page"] for hit in hits}, {1, 2})
        context, evidence = rag.pack_context(hits)
        self.assertIn("PDF 第 1 页", context)
        self.assertIn("PDF 第 2 页", context)
        for hit in evidence:
            self.assertEqual(hit["location"], hit["metadata"]["location"])
            if "定制商品" in hit["text"]:
                self.assertEqual(hit["metadata"]["page"], 2)
        generated = rag.ask(pipeline, "会员", top_k=100,
                            generator=lambda q, c: {"answer": "测试替身", "model": "test-only"})
        self.assertEqual(generated["evidence"], evidence)
        path.unlink()
        with pymupdf.open() as doc:
            doc.new_page()
            doc.save(path)
        with self.assertRaises(ValueError):
            pipeline.ingest(path)
        self.assertEqual(initial, set(pipeline.collection.get(include=[])["ids"]))


if __name__ == "__main__":
    unittest.main()
