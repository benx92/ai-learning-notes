from pathlib import Path
from openpyxl import load_workbook


def read_excel_units(path, sheet_name, header_row=1):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = next(sheet.iter_rows(min_row=header_row, max_row=header_row,
                                       values_only=True))
        if not any(value is not None and value != "" for value in headers):
            raise ValueError("指定行没有表头，请核对位置")
        for number, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                header_row + 1):
            pairs = []
            for i, value in enumerate(values):
                if value is None or value == "":
                    continue
                field = headers[i]
                if field is None or field == "":
                    field = f"列{i+1}"
                pairs.append(f"{field}={value}")
            if pairs:
                yield {"text": f"工作表={sheet_name}；" + "；".join(pairs),
                       "metadata": {"source": str(Path(path).resolve()),
                                    "sheet": sheet_name, "row": number}}
    finally:
        workbook.close()


if __name__ == "__main__":
    for unit in read_excel_units("policy.xlsx", "退货规定", header_row=1):
        print(unit)
