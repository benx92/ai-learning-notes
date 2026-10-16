from openpyxl import load_workbook


if __name__ == "__main__":
    formulas = load_workbook("policy.xlsx", data_only=False)
    cached = load_workbook("policy.xlsx", data_only=True)
    print("公式或原值：", formulas["退货规定"]["B2"].value)
    print("已保存的计算值：", cached["退货规定"]["B2"].value)
    formulas.close()
    cached.close()
