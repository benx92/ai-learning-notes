from pptx import Presentation


def read_slide_texts(shapes):
    for shape in shapes:
        if shape.has_text_frame and shape.text.strip():
            yield shape.text.strip()
        if shape.has_table:
            for row in shape.table.rows:
                yield " | ".join(cell.text.strip() for cell in row.cells)
        if hasattr(shape, "shapes"):
            yield from read_slide_texts(shape.shapes)


def read_slide_notes(slide):
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""



if __name__ == "__main__":
    for number, slide in enumerate(Presentation("policy.pptx").slides, 1):
        texts = list(read_slide_texts(slide.shapes))
        notes = read_slide_notes(slide)
        print(f"第{number}张画面文字：", "\n".join(texts) or "（未取得）")
        print(f"第{number}张备注：", notes or "（未取得）")
