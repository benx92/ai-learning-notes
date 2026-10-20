"""创建与文本案例规则相同的Word样例；拒绝覆盖已有文件。"""

import argparse
from pathlib import Path

from docx import Document


def build_sample(path):
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"保留已有文件：{path}")
    document = Document()
    document.add_paragraph("示例售后政策，仅用于学习。")
    for heading, rule in [
        ("会员订单", "签收后15日内可以申请退货。"),
        ("普通订单", "签收后7日内可以申请退货。"),
        ("定制商品例外", "定制商品不支持无理由退货。"),
        ("退款到账", "审核通过后3个工作日内到账。"),
    ]:
        document.add_paragraph(heading, style="Heading 1")
        document.add_paragraph(rule)
    document.save(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path,
                        default=Path(__file__).with_name("policy.docx"))
    print(build_sample(parser.parse_args().output))
